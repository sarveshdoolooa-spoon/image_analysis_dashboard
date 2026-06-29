"""
Car Plate Viewer — Streamlit edition.

Preserves the original Flask/HTML UI by embedding templates/index.html inside
st.components.v1.html(). Images and exports are served by a background Flask
thread (app.py routes) so the HTML payload stays small.
"""

from __future__ import annotations

import base64
import io
import json
import re
import socket
import threading
import time
from pathlib import Path

import openpyxl
import streamlit as st
import streamlit.components.v1 as components
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import (
    BASE_PATH,
    app as flask_app,
    build_journeys_py,
    get_image_folder,
    get_retrieval_folders,
    parse_excel,
)

APP_DIR = Path(__file__).parent
INDEX_HTML = APP_DIR / "templates" / "index.html"
API_HOST = "127.0.0.1"
API_PORT = 8765

HIDE_STREAMLIT = """
<style>
    #MainMenu, footer, header[data-testid="stHeader"] { visibility: hidden; height: 0; }
    [data-testid="stSidebar"], [data-testid="collapsedControl"] { display: none; }
    [data-testid="stAppViewContainer"] { padding: 0; }
    .block-container { padding: 0 !important; max-width: 100% !important; }
    iframe { border: none !important; width: 100% !important; }
    [data-testid="stVerticalBlock"] > div:has(iframe) { height: calc(100vh - 1rem); }
</style>
"""

IFRAME_CSS = """
    html, body {
        height: auto !important;
        min-height: 100% !important;
        overflow: auto !important;
        background: #0f1117 !important;
        color: #e2e8f0 !important;
    }
    body { display: flex; flex-direction: column; min-height: 900px; }
    main { flex: 1 1 auto; min-height: 400px; overflow: auto !important; }
"""


# ── Background Flask (images + export API from app.py) ───────────────────────


def _api_reachable() -> bool:
    try:
        with socket.create_connection((API_HOST, API_PORT), timeout=0.25):
            return True
    except OSError:
        return False


def _run_flask() -> None:
    flask_app.run(
        host=API_HOST,
        port=API_PORT,
        debug=False,
        use_reloader=False,
        threaded=True,
    )


@st.cache_resource(show_spinner=False)
def ensure_api_server() -> str:
    if not _api_reachable():
        thread = threading.Thread(target=_run_flask, daemon=True)
        thread.start()
        for _ in range(50):
            if _api_reachable():
                break
            time.sleep(0.1)
    if not _api_reachable():
        raise RuntimeError(f"Image/API server did not start on {API_HOST}:{API_PORT}")
    return f"http://{API_HOST}:{API_PORT}"


# ── Session defaults ──────────────────────────────────────────────────────────


def _init_session_state() -> None:
    defaults = {
        "folder": None,
        "view": "table",
        "search": "",
        "filter_type": "",
        "journey_filter": "all",
        "pm_threshold": 40,
        "dash_hide_no_images": False,
        "lightbox": None,
    }
    for key, val in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = val


# ── Data loading (same rules as app.py) ───────────────────────────────────────


def folder_label(name: str) -> str:
    ts = name.split("_")[-1] if "_" in name else ""
    if len(ts) == 14:
        return f"{ts[0:4]}-{ts[4:6]}-{ts[6:8]} {ts[8:10]}:{ts[10:12]}:{ts[12:14]}"
    return name


def list_folders_api() -> list[dict]:
    return [{"name": n, "label": folder_label(n)} for n in get_retrieval_folders()]


@st.cache_data(show_spinner=False)
def load_folder_records(folder_name: str) -> tuple[list[dict], dict | None, str | None]:
    if folder_name == "__ALL__":
        all_records: list[dict] = []
        for fn in get_retrieval_folders():
            data, err = parse_excel(fn)
            if err or not data:
                continue
            for r in data["records"]:
                r["folder"] = fn
            all_records.extend(data["records"])
        all_records.sort(key=lambda r: r.get("enter_time") or "")
        return all_records, {"title": "All Folders Combined"}, None

    data, err = parse_excel(folder_name)
    if err:
        return [], None, err
    for r in data["records"]:
        r.setdefault("folder", folder_name)
    return data["records"], data.get("meta"), None


# ── Excel export (same as app.py api_export_unmatched) ────────────────────────


def build_unmatched_excel(
    records: list[dict],
    folder: str,
    export_type: str,
    include_no_images: bool,
) -> tuple[bytes, str]:
    journeys = build_journeys_py(records)

    if export_type == "entry-only":
        rows = [
            {"rec": j["entry"], "unmatched_type": "Entry Without Exit"}
            for j in journeys if j["entry"] and not j["exit"]
        ]
    elif export_type == "exit-only":
        rows = [
            {"rec": j["exit"], "unmatched_type": "Exit Without Entry"}
            for j in journeys if not j["entry"] and j["exit"]
        ]
    else:
        rows = [
            {"rec": j["entry"], "unmatched_type": "Entry Without Exit"}
            for j in journeys if j["entry"] and not j["exit"]
        ] + [
            {"rec": j["exit"], "unmatched_type": "Exit Without Entry"}
            for j in journeys if not j["entry"] and j["exit"]
        ]

    if not include_no_images:
        rows = [item for item in rows if item["rec"].get("images")]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Unmatched Records"

    hdr_fill = PatternFill("solid", fgColor="1A1D27")
    hdr_font = Font(bold=True, color="E2E8F0")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        "Plate", "Unmatched Type", "Enter/Exit",
        "Enter Time", "Depart Time", "Duration",
        "Gate", "Region", "Result", "Reason",
        "Image 1", "Image 2", "Image 3", "Image 4",
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = hdr_align
    ws.row_dimensions[1].height = 30

    col_widths = [15, 22, 12, 22, 22, 12, 20, 14, 10, 40, 16, 16, 16, 16]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    IMG_W, IMG_H = 108, 78
    ROW_H_PT = 62
    data_align = Alignment(vertical="center", wrap_text=True)

    for row_idx, item in enumerate(rows, start=2):
        rec = item["rec"]
        ws.append([
            rec.get("plate", ""),
            item["unmatched_type"],
            rec.get("enter_exit", ""),
            rec.get("enter_time", ""),
            rec.get("depart_time", ""),
            rec.get("duration", ""),
            rec.get("gate", ""),
            rec.get("region", ""),
            rec.get("result", ""),
            rec.get("reason", ""),
        ])
        ws.row_dimensions[row_idx].height = ROW_H_PT
        for col_idx in range(1, 11):
            ws.cell(row_idx, col_idx).alignment = data_align

        type_cell = ws.cell(row_idx, 2)
        type_cell.font = Font(
            color="F59E0B" if item["unmatched_type"] == "Entry Without Exit" else "EF4444",
            bold=True,
        )

        rec_folder = rec.get("folder", folder if folder != "__ALL__" else "")
        img_folder = get_image_folder(rec_folder)
        for img_offset, fname in enumerate(rec.get("images", [])[:4]):
            img_path = img_folder / fname
            if not img_path.exists():
                continue
            try:
                xl_img = XLImage(str(img_path))
                xl_img.width = IMG_W
                xl_img.height = IMG_H
                col_letter = get_column_letter(11 + img_offset)
                ws.add_image(xl_img, f"{col_letter}{row_idx}")
            except Exception:
                pass

    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    type_label = {
        "entry-only": "entry_without_exit",
        "exit-only": "exit_without_entry",
    }.get(export_type, "all_unmatched")
    folder_label_str = folder[:30].replace(" ", "_") if folder != "__ALL__" else "all_folders"
    filename = f"unmatched_{type_label}_{folder_label_str}.xlsx"
    return buf.getvalue(), filename


# ── HTML shell (CSS + JS from index.html, adapted for Streamlit) ──────────────


def _extract_style_block() -> str:
    text = INDEX_HTML.read_text(encoding="utf-8")
    m = re.search(r"<style>(.*?)</style>", text, re.DOTALL)
    return m.group(1) if m else ""


def _extract_script_block() -> str:
    text = INDEX_HTML.read_text(encoding="utf-8")
    m = re.search(r"<script>(.*?)</script>", text, re.DOTALL)
    return m.group(1) if m else ""


def _patch_js(js: str, api_base: str) -> str:
    """Adapt the original client script for embedded Streamlit data."""
    js = js.replace(
        "let allRecords = [], filtered = [], currentView = 'table', currentFolder = '';",
        "let allRecords = [], filtered = [], currentView = 'table';",
    )

    js = js.replace(
        "const skipType = currentView === 'journey' || currentView === 'dashboard';",
        "const skipType = currentView === 'journey' || currentView === 'dashboard' || currentView === 'possiblematches';",
    )

    init_start = js.index("async function init()")
    load_start = js.index("async function loadFolder()")
    js = (
        js[:init_start]
        + """function init() {
    const sel = document.getElementById('folderSelect');
    sel.innerHTML = '';
    const allOpt = document.createElement('option');
    allOpt.value = '__ALL__';
    allOpt.textContent = '★ All Folders Combined';
    sel.appendChild(allOpt);
    foldersList.forEach(f => {
      const o = document.createElement('option');
      o.value = f.name;
      o.textContent = f.label + '  (' + f.name.replace('Entrance and Exit Retrieval_','') + ')';
      sel.appendChild(o);
    });
    sel.value = currentFolder;
    allRecords = STREAMLIT_RECORDS.slice();
    journeyFilter = 'all';
    pmThreshold = STREAMLIT_PM_THRESHOLD;
    dashHideNoImages = STREAMLIT_DASH_HIDE_NO_IMAGES;
    currentView = STREAMLIT_VIEW;
    ['table','grid','journey','dashboard','possiblematches'].forEach(id => {
      const btn = document.getElementById('btn' + id.charAt(0).toUpperCase() + id.slice(1));
      if (btn) btn.classList.toggle('active', id === currentView);
    });
    const noFilter = currentView === 'journey' || currentView === 'dashboard' || currentView === 'possiblematches';
    document.getElementById('filterType').disabled = noFilter;
    document.getElementById('searchInput').value = STREAMLIT_SEARCH;
    document.getElementById('filterType').value = STREAMLIT_FILTER_TYPE;
    if (STREAMLIT_JOURNEY_FILTER !== 'all') journeyFilter = STREAMLIT_JOURNEY_FILTER;
    applyFilter();
  }

"""
        + js[load_start:]
    )

    load_start = js.index("async function loadFolder()")
    load_end = js.index("  /* ── Filter", load_start)
    js = js[:load_start] + """async function loadFolder() {
    const folder = document.getElementById('folderSelect').value;
    if (!folder) return;
    currentFolder = folder;
    document.getElementById('main').innerHTML =
      `<div class="state-msg"><div class="spinner"></div><p>Loading records…</p></div>`;

    const url = folder === '__ALL__'
      ? `${API_BASE}/api/all-records`
      : `${API_BASE}/api/records/${encodeURIComponent(folder)}`;
    try {
      const res = await fetch(url);
      if (!res.ok) {
        document.getElementById('main').innerHTML =
          `<div class="state-msg"><div class="icon">⚠️</div><p>Failed to load records.</p></div>`;
        return;
      }
      allRecords = (await res.json()).records || [];
      journeyFilter = 'all';
      applyFilter();
    } catch (_) {
      document.getElementById('main').innerHTML =
        `<div class="state-msg"><div class="icon">⚠️</div><p>Failed to load records.</p></div>`;
    }
  }

""" + js[load_end:]

    api = api_base.rstrip("/")
    js = js.replace(
        """function imgUrl(filename, folder) {
    const f = folder || currentFolder;
    return `/images/${encodeURIComponent(f)}/${encodeURIComponent(filename)}`;
  }""",
        f"""function imgUrl(filename, folder) {{
    const f = folder || currentFolder;
    return `{api}/images/${{encodeURIComponent(f)}}/${{encodeURIComponent(filename)}}`;
  }}""",
    )

    js = js.replace(
        """function exportUnmatched(type) {
    const params = new URLSearchParams({
      folder: currentFolder,
      type,
      include_no_images: dashHideNoImages ? '0' : '1',
    });
    window.location.href = `/api/export-unmatched?${params}`;
  }""",
        f"""function exportUnmatched(type) {{
    const params = new URLSearchParams({{
      folder: currentFolder,
      type,
      include_no_images: dashHideNoImages ? '0' : '1',
    }});
    window.open(`{api}/api/export-unmatched?${{params}}`, '_blank');
  }}""",
    )

    js = js.replace("  init();", "\n  init();")
    return js


def build_viewer_html(
    records: list[dict],
    folder: str,
    folders: list[dict],
    api_base: str,
) -> str:
    css = _extract_style_block()
    js = _patch_js(_extract_script_block(), api_base)

    # Inject server data BEFORE the let declarations (no duplicate const names).
    bootstrap = {
        "API_BASE": api_base.rstrip("/"),
        "STREAMLIT_RECORDS": records,
        "foldersList": folders,
        "STREAMLIT_VIEW": st.session_state.view,
        "STREAMLIT_SEARCH": st.session_state.search,
        "STREAMLIT_FILTER_TYPE": st.session_state.filter_type,
        "STREAMLIT_JOURNEY_FILTER": st.session_state.journey_filter,
        "STREAMLIT_PM_THRESHOLD": st.session_state.pm_threshold,
        "STREAMLIT_DASH_HIDE_NO_IMAGES": st.session_state.dash_hide_no_images,
    }
    bootstrap_js = f"const {', '.join(f'{k} = {json.dumps(v)}' for k, v in bootstrap.items())};\n"

    body = INDEX_HTML.read_text(encoding="utf-8")
    body_start = body.index("<body>")
    body_end = body.index("<!-- Lightbox -->")
    body_fragment = body[body_start:body_end]

    lightbox_start = body.index("<!-- Lightbox -->")
    lightbox_end = body.index("<script>")
    lightbox_fragment = body[lightbox_start:lightbox_end]

    pre_js = f"let currentFolder = {json.dumps(folder)};\n"

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>Car Plate Viewer — SoFlo1543</title>
  <script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.3/dist/chart.umd.min.js"></script>
  <style>{css}{IFRAME_CSS}</style>
</head>
{body_fragment}
{lightbox_fragment}
<script>
{bootstrap_js}
{pre_js}
{js}
</script>
</body>
</html>"""


# ── Query-param ↔ session_state sync ──────────────────────────────────────────


def _sync_state_from_query(folders: list[str]) -> None:
    qp = st.query_params
    if "folder" in qp:
        st.session_state.folder = qp["folder"]
    if st.session_state.folder is None and folders:
        st.session_state.folder = folders[0]


# ── Main ──────────────────────────────────────────────────────────────────────


def main() -> None:
    st.set_page_config(
        page_title="Car Plate Viewer — SoFlo1543",
        page_icon="🌍",
        layout="wide",
        initial_sidebar_state="collapsed",
    )
    _init_session_state()
    st.markdown(HIDE_STREAMLIT, unsafe_allow_html=True)

    folders_raw = get_retrieval_folders()
    if not folders_raw:
        st.error(f"No retrieval folders found under `{BASE_PATH}`.")
        st.stop()

    try:
        api_base = ensure_api_server()
    except RuntimeError as exc:
        st.error(str(exc))
        st.stop()

    folders_api = list_folders_api()
    _sync_state_from_query(folders_raw)

    folder = st.session_state.folder
    if folder not in folders_raw and folder != "__ALL__":
        folder = folders_raw[0]
        st.session_state.folder = folder

    with st.spinner("Loading records…"):
        records, _meta, err = load_folder_records(folder)

    if err:
        st.error(err)
        st.stop()

    html = build_viewer_html(
        records=records,
        folder=folder,
        folders=folders_api,
        api_base=api_base,
    )

    components.html(html, height=900, scrolling=True)

    with st.expander("Export unmatched records (Streamlit)", expanded=False):
        include_imgs = not st.session_state.dash_hide_no_images
        for label, etype in (
            ("Entry without exit", "entry-only"),
            ("Exit without entry", "exit-only"),
            ("All unmatched", "both"),
        ):
            data, fname = build_unmatched_excel(records, folder, etype, include_imgs)
            st.download_button(
                label,
                data,
                fname,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=f"dl_{etype}",
            )


if __name__ == "__main__":
    main()
